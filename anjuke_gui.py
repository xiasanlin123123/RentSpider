# -*- coding: utf-8 -*-
"""
安居客郑州租房爬虫（图形界面版）

用法：选大区 -> 选商圈 -> 点「开始采集」-> 完成后 Excel 保存在桌面上。
遇到人机验证会自动打开浏览器（Chrome / Edge 都支持），
在浏览器里完成验证后点「确定」继续。
"""
import os
import re
import sys
import time
import random
import threading
import webbrowser
import subprocess

import requests
import parsel
from tkinter import Tk, StringVar, DISABLED, NORMAL
from tkinter import ttk, messagebox

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except Exception:                                   # 极少见：缺库时降级为 CSV
    HAS_XLSX = False
    import csv

APP_TITLE = '安居客郑州租房爬虫'

# 请求头里的 UA。只用来看起来像浏览器，和用户实际用 Chrome 还是 Edge 无关
# （服务端是 IP 级限流，实测换 UA 没有任何影响）。选了具体浏览器会同步成它的 UA。
UA_BASE = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
           '(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36')
UA_CHROME = UA_BASE
UA_EDGE = UA_BASE + ' Edg/140.0.0.0'
UA = UA_EDGE                                        # 默认

# ---------------------------------------------------------------- 地区数据
# 一级地区 -> (区域 URL 片段, ['商圈名:商圈 URL 片段', ...])
# 数据来源：2026-10-01 从安居客郑州租房官网「区域」筛选栏逐个抓取。
# 注意：商圈的 URL 片段（slug）不能按拼音猜——官网用的是带随机后缀的内部编码
# （例如「博颂路」是 jinshui-q-renminluzz，不是 jinshui-q-baosonglu）。
# 猜的 slug 不会报错，而是静默兜底成整个大区，非常容易误判，所以这里全部照抄官网。
REGIONS = {
    '全部（郑州所有地区）': ('', []),
    '金水': ('jinshui', [
        '博颂路:jinshui-q-renminluzz', '北环路:jinshui-q-beihuanlu', '陈寨:jinshui-q-chenz',
        '东风路:jinshui-q-dongfengluzz', '大石桥:jinshui-q-dashiqiaozz', '东明路:jinshui-q-xinliulu',
        '东三街:jinshui-q-dongsanjie', '21世纪社区:jinshui-q-sjsq', '丰产路:jinshui-q-xihanzhai',
        '丰庆路:jinshui-q-fengqinglu', '丰乐路:jinshui-q-fengll', '福彩路:jinshui-q-fucail',
        '国基路:jinshui-q-guojilu', '广电南路:jinshui-q-zhongfangyuan', '花园路:jinshui-q-huayuanlu',
        '黄河路:jinshui-q-huangheluzz', '红专路:jinshui-q-hzl', '经七路:jinshui-q-jingbalu',
        '金水路:jinshui-q-beilinlu', '经三路:jinshui-q-fenghuangtaizz', '金水周边:jinshui-q-jinshuiqu',
        '科技市场:jinshui-q-kjsc', '绿荫广场:jinshui-q-lygc', '曼哈顿:jinshui-q-mhdzz',
        '南阳路:jinshui-q-xishakou', '农业路:jinshui-q-nongyelu', '农科路:jinshui-q-nongkel',
        '三全路:jinshui-q-jinchengzz', '索凌路:jinshui-q-suolinglu', '沙口路:jinshui-q-shakl',
        '水上公园:jinshui-q-ssgy', '省政府:jinshui-q-szf', '天明路:jinshui-q-tianml',
        '文化路:jinshui-q-dapu', '未来路:jinshui-q-yanzhuang', '文博广场:jinshui-q-wbgc',
        '新通桥:jinshui-q-xtq', '鑫苑路:jinshui-q-xinyl', '英协路:jinshui-q-yingxielu',
        '燕庄:jinshui-q-yzhuang', '玉凤路:jinshui-q-yfl', '郑州东站:jinshui-q-dongzhanpian',
        '政七街:jinshui-q-duling', '中州大道:jinshui-q-longzihuzz', '中医院:jinshui-q-zyy',
    ]),
    '二七': ('erqic', [
        '碧云路:erqic-q-byl', '长江路:erqic-q-changjiangluzz', '大学路:erqic-q-daxuelu',
        '大学南路:erqic-q-dxnl', '二七周边:erqic-q-erqiqu', '古玩城:erqic-q-gwczz',
        '淮河路:erqic-q-luzhai', '火车站:erqic-q-dehuajie', '华中片:erqic-q-huazhp',
        '河医片:erqic-q-heyip', '淮南街:erqic-q-hnj', '淮北街:erqic-q-hbj',
        '航海路:erqic-q-hhlzz', '解放路:erqic-q-jiefanglu', '京广路:erqic-q-jingguanglu',
        '交通路:erqic-q-jtl', '连云路:erqic-q-lyunlu', '铭功路:erqic-q-minggonglu',
        '棉纺东路:erqic-q-mfdl', '南三环:erqic-q-nshzz', '庆丰街:erqic-q-qingfj',
        '嵩山路:erqic-q-qiliyan', '嵩山南路:erqic-q-ssnl', '升龙国际:erqic-q-slgj',
        '桃源路:erqic-q-tyl', '万达广场:erqic-q-wandgc', '兴华南街:erqic-q-xinghuanj',
        '一马路:erqic-q-yimalu', '政通路:erqic-q-ztl',
    ]),
    '郑东新区': ('zhengdongxinqu', [
        '北龙湖:zhengdongxinqu-q-beilonghuzz', 'CBD:zhengdongxinqu-q-cbdneihuanlu', '东风东路:zhengdongxinqu-q-dongfengdonglu',
        '东风南路:zhengdongxinqu-q-dongfengnanlu', '鸿园片区:zhengdongxinqu-q-hongyuanpianquabc', '黄河东路:zhengdongxinqu-q-huanghedonglu',
        '黄河南路:zhengdongxinqu-q-huanghenanlu', '祭城:zhengdongxinqu-q-jczz', '聚源路:zhengdongxinqu-q-juyuanl',
        '金水东路:zhengdongxinqu-q-jinshuidonglu', '九如路:zhengdongxinqu-q-jrl', '康宁街:zhengdongxinqu-q-kangnignjie',
        '康平路:zhengdongxinqu-q-kangpinglu', '龙子湖:zhengdongxinqu-q-lzhzz', '绿地老街:zhengdongxinqu-q-lvdilaojie',
        '绿城百合:zhengdongxinqu-q-lvchengbh', '列里路:zhengdongxinqu-q-lielilu', '农业东路:zhengdongxinqu-q-nonglindonglu',
        '农业南路:zhengdongxinqu-q-nongynanlu', '七里河:zhengdongxinqu-q-qilh', '商都路:zhengdongxinqu-q-shangdl',
        '商鼎路:zhengdongxinqu-q-shangdingl', '通泰路:zhengdongxinqu-q-tongtailu', '天赋路:zhengdongxinqu-q-tfl',
        '祥盛街:zhengdongxinqu-q-xiangshengjie', '心怡路:zhengdongxinqu-q-xinyilu', '兴荣街:zhengdongxinqu-q-xrj',
        '新东站片:zhengdongxinqu-q-xdzp', '杨金片区:zhengdongxinqu-q-yangjinpianquabc', '永平路:zhengdongxinqu-q-yongpinglu',
        '众意西路:zhengdongxinqu-q-zhongyixilu', '郑东新区周边:zhengdongxinqu-q-zhengdongxqzb',
    ]),
    '中原': ('zhongyuanb', [
        '碧沙岗:zhongyuanb-q-bishagang', '常西湖:zhongyuanb-q-changxihus', '帝湖:zhongyuanb-q-jiansheluzz',
        '航海西路:zhongyuanb-q-shiyangsi', '淮河西路:zhongyuanb-q-huaihexilu', '建设西路:zhongyuanb-q-jianshexilu',
        '凯旋门:zhongyuanb-q-kaixuanmen', '林山寨:zhongyuanb-q-linshanzhai', '陇海西路:zhongyuanb-q-longhaixilu',
        '棉纺路:zhongyuanb-q-mianfanglu', '秦岭路:zhongyuanb-q-qinlinglu', '汝河路:zhongyuanb-q-ruhelu',
        '三官庙:zhongyuanb-q-sanguanmiao', '嵩山北路:zhongyuanb-q-songshanbeilu', '柿园:zhongyuanb-q-shiy',
        '石佛镇:zhongyuanb-q-shifz', '桐柏路:zhongyuanb-q-zhouxinzhuang', '桐柏北路:zhongyuanb-q-tbbl',
        '桐柏南路:zhongyuanb-q-tbnl', '五龙口:zhongyuanb-q-lvdongcun', '五一公园:zhongyuanb-q-wuyigy',
        '文化宫路:zhongyuanb-q-wenhuagl', '伊河路:zhongyuanb-q-yihelu', '中原西路:zhongyuanb-q-zhongyuanxilu',
        '郑州市一中:zhongyuanb-q-zzsyz', '郑上路小学:zhongyuanb-q-zslxx', '中原周边:zhongyuanb-q-zhongyuanqu',
    ]),
    '管城': ('guanchenga', [
        '安徐庄:guanchenga-q-axzzz', '城东路:guanchenga-q-chengdonglu', '东大街:guanchenga-q-dongdajiezz',
        '东明路南段:guanchenga-q-dongml', '东太康路:guanchenga-q-dongtaikanglu', '二里岗:guanchenga-q-erligang',
        '凤台路:guanchenga-q-fengtailuzz', '富田太阳城:guanchenga-q-fttyczz', '管南片:guanchenga-q-guannanpian',
        '管城周边:guanchenga-q-guanchengqu', '航海东路:guanchenga-q-shangyingjie', '陇海东路:guanchenga-q-xishilipu',
        '南关:guanchenga-q-nanguanzz', '南大街:guanchenga-q-ndj', '人民路:guanchenga-q-zzrenminlu',
        '商代遗址:guanchenga-q-shangdaiyizhi', '商城路:guanchenga-q-shangcl', '世纪欢乐园:guanchenga-q-shihlc',
        '西大街:guanchenga-q-xidajiezz', '紫荆山路:guanchenga-q-zijingshannanlu', '郑汴路:guanchenga-q-zhengbianlu',
    ]),
    '新郑': ('xinzhengshi', [
        '龙湖双湖大道:xinzhengshi-q-longhushuanghudadao', '龙湖阳光大道:xinzhengshi-q-longhuyangguangdadao', '龙湖沙窝李:xinzhengshi-q-longhushawoli',
        '龙湖华南城:xinzhengshi-q-longhuhuanancheng', '新郑城区:xinzhengshi-q-xinzheng',
    ]),
    '高新区': ('gaoxinqub', [
        '大学科技园西区:gaoxinqub-q-daxuekejiyuanxiqu', '公园道:gaoxinqub-q-gongyuandaoabc', '科学大道:gaoxinqub-q-kexuedadao',
        '莲花街:gaoxinqub-q-lainhuajie', '瑞达路:gaoxinqub-q-ruidalu',
    ]),
    '惠济': ('huiji', [
        '北大学城:huiji-q-beidaxuecheng', '大河路:huiji-q-dahelu', '江山路:huiji-q-jiangshanl',
        '开元路:huiji-q-kaiyuanl', '刘寨:huiji-q-liuzhai', '邙山片:huiji-q-mangshanpian',
        '省体育中心:huiji-q-xinchengzz', '思念果岭:huiji-q-huijiqu', '迎宾路:huiji-q-yingbinlu',
        '英才街:huiji-q-yingcaij',
    ]),
    '中牟': ('zhongmou', [
        '白沙镇:zhongmou-q-bszzz', '百乐汇购物中心:zhongmou-q-blhgwzxzz', '杉杉奥特莱斯:zhongmou-q-ssatlszz',
        '中牟城区:zhongmou-q-shangjiequ', '中牟县人民医院:zhongmou-q-zhongmuxianrenminyiyuanaa',
    ]),
    '荥阳': ('zzxingyangshi', [
        '洞林湖:zzxingyangshi-q-dlhzz', '荥阳城区:zzxingyangshi-q-yingyang', '忆江南:zzxingyangshi-q-yijiangnan',
    ]),
    '经开区': ('jingkaiz', [
        '滨河新区:jingkaiz-q-bhxqzz', '第三大街:jingkaiz-q-disandj', '第八大街:jingkaiz-q-dibadj',
        '经开物流园区:jingkaiz-q-jkqwlyq', '远大理想城:jingkaiz-q-yuandalxc',
    ]),
    '航空港': ('hangkongganga', [
        '北港:hangkongganga-q-beigangzz', '南港:hangkongganga-q-nangangzz', '薛店:hangkongganga-q-xuediana',
    ]),
    '上街': ('zzshangjiequ', [
        '工业路:zzshangjiequ-q-gongyelu', '济源路:zzshangjiequ-q-jiyuanluzz', '矿山:zzshangjiequ-q-kuangshan',
        '上街周边:zzshangjiequ-q-shangjiezhoubian', '新安路:zzshangjiequ-q-xinanluzz', '中心路:zzshangjiequ-q-zhongxinlu',
    ]),
    '新密': ('xinmishi', [
        '曲梁镇:xinmishi-q-quliangz', '新密城区:xinmishi-q-xinmichenqu',
    ]),
    '巩义': ('gongyishi', [
        '巩义城区:gongyishi-q-gongyichenqu',
    ]),
    '登封': ('dengfengshi', [
        '登封城区:dengfengshi-q-dengfengchenqu',
    ]),
    '郑州周边': ('zhengzhouzhoubian', [
        '新乡:zhengzhouzhoubian-q-xinxiangzz',
    ]),
}

SITE = 'https://zz.zu.anjuke.com/fangyuan/'
BIZ_ALL = '全部（不限商圈）'                         # 二级下拉第一项：只要整个大区

REGION_LIST = list(REGIONS)                         # 保序的大区名
REGION_BIZ = {r: [b.split(':', 1)[0] for b in v[1]] for r, v in REGIONS.items()}
REGION_BIZ_SLUG = {r: dict(b.split(':', 1) for b in v[1])
                   for r, v in REGIONS.items()}

FIELDS = ['描述', '几室几厅', '价格', '楼层', '是否合租',
          '朝向', '电梯', '地铁', '地址', '房源链接']

# 表头 / 蓝白相间 / 边框 的配色（8 位 ARGB，兼容性最好）
HEAD_BG = 'FF1F4E79'        # 表头深蓝
ZEBRA_BG = 'FFDCE9F7'       # 隔行浅蓝
GRID_BG = 'FFBFD4EA'        # 单元格边框
FONT_NAME = '微软雅黑'
COL_WIDTH = [42, 12, 11, 13, 10, 8, 8, 16, 34, 13]

PAGE_GAP = (3.0, 5.0)       # 每页之间的随机间隔（秒）
MAX_BLOCK = 5               # 连续这么多次过不了验证就放弃，避免无限弹窗
SAVE_EVERY = 3              # 每采集这么多页就落盘一次，中途中断不丢数据

DEFAULT_BROWSER = '系统默认浏览器'


# ---------------------------------------------------------------- 工具函数

def base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def resource(rel):
    base = getattr(sys, '_MEIPASS', base_dir())
    return os.path.join(base, rel)


def desktop_dir():
    """取桌面真实路径（兼容 OneDrive 重定向），失败则退回程序所在目录。"""
    try:
        import winreg
        key = winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            r'Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders')
        p = winreg.QueryValueEx(key, 'Desktop')[0]
        if os.path.isdir(p):
            return p
    except Exception:
        pass
    for cand in (os.path.join(os.path.expanduser('~'), 'Desktop'),
                 os.path.join(os.path.expanduser('~'), '桌面')):
        if os.path.isdir(cand):
            return cand
    return base_dir()


def resolve_target(region, biz):
    """把「大区 + 商圈」翻译成列表页地址和文件名标签。

    返回 (列表页 URL, 标签)。选了具体商圈用商圈 URL，否则用大区 URL。
    """
    slug, _ = REGIONS[region]
    tag = region.split('（')[0]
    if biz and biz != BIZ_ALL:
        bslug = REGION_BIZ_SLUG.get(region, {}).get(biz)
        if bslug:
            return f'{SITE}{bslug}/', f'{tag}-{biz}'
    if slug:
        return f'{SITE}{slug}/', tag
    return SITE, '全部'


# ---- 浏览器检测 ------------------------------------------------------------
# 程序本身和浏览器内核无关：请求用 requests 发出，浏览器只用来过验证码。
# 无论用户平时用 Chrome 还是 Edge 都能正常跑，这里只是让用户指定
# 「验证页面在哪个浏览器里打开」。

def detect_browsers():
    """返回本机已安装的浏览器 [(显示名, exe 完整路径), ...]。"""
    import winreg
    app_paths = r'SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths'
    cands = [
        ('Google Chrome', 'chrome.exe',
         [r'C:\Program Files\Google\Chrome\Application\chrome.exe',
          r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
          r'%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe']),
        ('Microsoft Edge', 'msedge.exe',
         [r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
          r'C:\Program Files\Microsoft\Edge\Application\msedge.exe']),
    ]
    found = []
    for name, exe, paths in cands:
        path = ''
        for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
            try:
                k = winreg.OpenKey(hive, f'{app_paths}\\{exe}')
                path = winreg.QueryValueEx(k, '')[0]
                winreg.CloseKey(k)
                if path:
                    break
            except Exception:
                continue
        if not (path and os.path.isfile(path)):     # 注册表没有再查常见路径
            for p in paths:
                p = os.path.expandvars(p)
                if os.path.isfile(p):
                    path = p
                    break
        if path and os.path.isfile(path):
            found.append((name, path))
    return found


def open_in_browser(browser_name, exe_path, url):
    """用指定浏览器打开网址；失败或未指定则交给系统默认浏览器。"""
    if exe_path:
        try:
            subprocess.Popen([exe_path, url], close_fds=True)
            return True
        except Exception:
            pass
    try:
        return webbrowser.open(url)
    except Exception:
        return False


def ua_for(browser_name):
    """让请求头 UA 和用户选的浏览器保持一致（一致性更好，非必需）。"""
    if browser_name and 'Chrome' in browser_name:
        return UA_CHROME
    return UA_EDGE


# ---- 人机验证识别 ----------------------------------------------------------
# 被拦截时安居客会返回一个 700 多字节的“跳板页”：HTTP 200、地址栏不变、
# 正文里只有一个隐藏 div#@@xxzlGatewayUrl，里面是要跳转的验证页地址。
# 页面由 JS 跳转，所以 requests 拿到的一定是这个跳板页，特征非常稳定。
GATEWAY_ID = '@@xxzlGatewayUrl'
BLOCK_TEXTS = ('访问过于频繁', '请输入验证码', '验证码校验', '完成验证')


def extract_gateway(html):
    """从跳板页里取出真正的验证页地址，取不到就返回空串。"""
    m = re.search(r'id="@@xxzlGatewayUrl"[^>]*>(.*?)</div>', html, re.S)
    if not m:
        return ''
    url = m.group(1).strip().replace('&amp;', '&')
    return url if url.startswith('http') else ''


def is_blocked(resp):
    """是否被人机验证拦截。

    注意：被拦截时 HTTP 状态码依然是 200，地址也不会跳转，
    所以必须看正文里的跳板页特征，不能只看状态码。
    """
    if 'antibot' in resp.url or 'callback.58.com' in resp.url:
        return True
    text = resp.text
    if GATEWAY_ID in text or 'callback.58.com' in text or '/antibot/' in text:
        return True
    return any(k in text for k in BLOCK_TEXTS)


def pick_tag(tags, keywords):
    for t in tags:
        for kw in keywords:
            if kw in t:
                return t
    return ''


def parse_page(html):
    """解析列表页 HTML，返回房源字典列表。"""
    rows = []
    for div in parsel.Selector(html).css('.zu-info'):
        try:
            title = div.css('h3 b.strongbox::text').get() or ''

            nums = div.css('p.details-item.tag b.strongbox::text').getall()
            num = f'{nums[0]}室{nums[1]}厅{nums[2]}平米' if len(nums) >= 3 else ''

            price = div.xpath('./following-sibling::div[contains(@class,"zu-side")]//strong[contains(@class,"price")]/text()').get()
            unit = div.xpath('./following-sibling::div[contains(@class,"zu-side")]//span[contains(@class,"unit")]/text()').get()

            # 地址：要用 string(.) 取整个节点（含 <a> 小区名和 <span> 分隔符）的文本。
            # 旧的 ::text 只拿直接文本节点，会把开头的「小区名」整段丢掉，
            # 结果只剩「中牟-杉杉奥特莱斯-仁安路」这半截，看起来像少了信息。
            addr_raw = div.css('address.details-item.tag').xpath('string(.)').get() or ''
            address_text = ' '.join(addr_raw.split())

            words = div.css('p.details-item.tag::text').getall()
            floor = words[4].strip() if len(words) > 4 else ''

            tags = div.css('p.details-item.bot-tag span.cls-common::text').getall()

            # 房源链接：普通房源是 /fangyuan/{id}，
            # 「安选」房源是 /gfangyuan/{id}（郑州中牟从第 30 页起大量出现）。
            # 旧写法只认 /fangyuan/，导致后半段房源的链接整片丢失。
            # 这里保留前缀原样，并丢掉 ?isauction=..&psid=.. 这类跟踪参数。
            href = div.css('h3 a::attr(href)').get() or ''
            m = re.search(r'/(g?fangyuan)/(\d+)', href)
            link = (f'https://zz.zu.anjuke.com/{m.group(1)}/{m.group(2)}'
                    if m else '')

            rows.append({
                '描述': title,
                '几室几厅': num,
                '价格': f'{price}{unit}' if price else '',
                '楼层': floor,
                '是否合租': pick_tag(tags, ['合租', '整租', '独立']),
                '朝向': pick_tag(tags, ['朝']),
                '电梯': pick_tag(tags, ['电梯']),
                '地铁': pick_tag(tags, ['号线', '地铁']),
                '地址': address_text,
                '房源链接': link,
            })
        except Exception:
            continue
    return rows


# ---- 输出：带样式的 Excel --------------------------------------------------

def save_xlsx(path, rows, sheet_name):
    """蓝白相间 + 表头冻结 + 链接可直接点击的 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or '房源')[:31]

    thin = Side(style='thin', color=GRID_BG)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name=FONT_NAME, size=10, bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor=HEAD_BG)
    zebra_fill = PatternFill('solid', fgColor=ZEBRA_BG)
    body_font = Font(name=FONT_NAME, size=10)
    link_font = Font(name=FONT_NAME, size=10, color='0563C1', underline='single')

    for c, name in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    for i, row in enumerate(rows):
        r = i + 2
        for c, name in enumerate(FIELDS, 1):
            v = row.get(name, '')
            cell = ws.cell(row=r, column=c)
            cell.border = border
            if name == '房源链接' and v:
                cell.value = '打开房源'
                cell.hyperlink = v
                cell.font = link_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.value = v
                cell.font = body_font
                cell.alignment = Alignment(
                    horizontal='left' if name in ('描述', '地址') else 'center',
                    vertical='center')
            if r % 2 == 0:
                cell.fill = zebra_fill

    for c, w in enumerate(COL_WIDTH, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    # 冻结「表头行 + 首列(描述)」：
    # 窗口较窄时，点最右边的「房源链接」会让整表向左横滚，把描述列开头几个字挡在屏幕外，
    # 看起来像「数据缺了一段」。冻住首列后，无论怎么横滚，描述都完整可见。
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(FIELDS))}{len(rows) + 1}'

    tmp = path + '.tmp'
    wb.save(tmp)
    os.replace(tmp, path)           # 原子替换，避免写一半被打断


def save_csv(path, rows):
    with open(path, 'w', encoding='utf-8-sig', newline='') as fi:
        w = csv.DictWriter(fi, fieldnames=FIELDS)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def save_table(path, rows, sheet_name):
    """落盘；文件被占用时自动换个名字，保证数据不丢。"""
    for n in range(6):
        target = path if n == 0 else path[:-5] + f'({n}).xlsx'
        try:
            if HAS_XLSX:
                save_xlsx(target, rows, sheet_name)
            else:
                save_csv(target[:-5] + '.csv', rows)
            return target
        except PermissionError:
            continue
        except Exception:
            break
    if HAS_XLSX:                    # Excel 写不出来就退回 CSV，至少数据在
        fallback = path[:-5] + '.csv'
        save_csv(fallback, rows)
        return fallback
    return path


# ---------------------------------------------------------------- 采集线程

class Crawler(threading.Thread):
    def __init__(self, app, url, tag, browser_name, browser_exe):
        super().__init__(daemon=True)
        self.app = app
        self.url = url                            # 列表页首页地址
        self.tag = tag                            # 只用于文件名 / sheet 名
        self.browser_name = browser_name
        self.browser_exe = browser_exe
        self.stop_event = threading.Event()

    def run(self):
        session = requests.Session()
        session.headers.update({'user-agent': ua_for(self.browser_name)})

        path = os.path.join(
            desktop_dir(), f'安居客_{self.tag}_{time.strftime("%Y%m%d_%H%M%S")}.xlsx')

        # 先访问一次首页拿到正常 Cookie，让请求看起来更像真人
        try:
            session.get('https://zz.zu.anjuke.com/', timeout=20)
        except Exception:
            pass

        rows_all = []
        seen = set()
        total = 0
        page = 1
        blocked_times = 0
        saved_path = path

        while not self.stop_event.is_set():
            url = self.url if page == 1 else f'{self.url}p{page}/'
            self.app.set_status(f'正在采集第 {page} 页…')

            try:
                resp = session.get(url, timeout=20)
            except Exception as e:
                self.app.set_status(f'网络异常，3 秒后重试：{str(e)[:40]}')
                time.sleep(3)
                continue

            # ---- 撞上人机验证 ----
            if is_blocked(resp):
                blocked_times += 1
                if blocked_times > MAX_BLOCK:
                    self.app.set_status('连续多次未通过验证，已停止')
                    self.app.notify(
                        '验证未通过',
                        '连续多次都没能通过安居客的人机验证。\n\n'
                        '建议：\n'
                        '1. 先点【停止】，等 5~10 分钟再试\n'
                        '2. 或者换一个网络环境（比如手机热点）\n\n'
                        '已采集到的数据会保留。')
                    break

                gate = extract_gateway(resp.text) or self.url
                if blocked_times == 1:
                    self.app.set_status('被人机验证拦截，已打开浏览器，请完成验证…')
                else:
                    self.app.set_status(f'仍未通过验证（第 {blocked_times} 次）…')

                if not self.app.ask_verify(gate, self.url, self.browser_name,
                                           self.browser_exe, first=(blocked_times == 1)):
                    break
                time.sleep(min(3 * blocked_times, 15))     # 退避
                continue

            blocked_times = 0

            rows = parse_page(resp.text)
            if not rows and page == 1:
                # 首屏就没数据，多半还是被拦（换了一种拦截形态），再试两次
                if blocked_times == 0 and getattr(self, '_empty_retry', 0) < 2:
                    self._empty_retry = getattr(self, '_empty_retry', 0) + 1
                    self.app.set_status('第 1 页暂时没拿到数据，重试中…')
                    time.sleep(3)
                    continue
                self.app.set_status('第 1 页没有数据，采集结束')
                break
            if not rows:
                self.app.set_status(f'第 {page} 页没有更多房源了，采集结束')
                break

            for row in rows:
                if self.stop_event.is_set():
                    break
                link = row['房源链接']
                if link:
                    if link in seen:
                        continue
                    seen.add(link)
                rows_all.append(row)
                total += 1

            self.app.set_count(total)
            if page % SAVE_EVERY == 1:      # 定期落盘
                saved_path = save_table(path, rows_all, self.tag)

            page += 1
            for _ in range(int(random.uniform(*PAGE_GAP) * 10)):   # 降速
                if self.stop_event.is_set():
                    break
                time.sleep(0.1)

        if rows_all:
            saved_path = save_table(path, rows_all, self.tag)
        self.app.finish(total, saved_path, self.stop_event.is_set())


# ---------------------------------------------------------------- 界面

class App:
    def __init__(self):
        self.root = Tk()
        self.root.title(APP_TITLE)
        self.root.resizable(False, False)
        try:
            self.root.iconbitmap(resource('app.ico'))
        except Exception:
            pass

        self.verify_done = threading.Event()
        self.verify_ok = False
        self.crawler = None

        # 本机装了哪些浏览器（Chrome / Edge）
        self.browsers = detect_browsers()
        self.browser_map = {DEFAULT_BROWSER: ''}
        for name, exe in self.browsers:
            self.browser_map[name] = exe

        first_region = REGION_LIST[0]
        self.region = StringVar(value=first_region)
        self.biz = StringVar(value=BIZ_ALL)
        self.browser = StringVar(value=DEFAULT_BROWSER)
        self.status = StringVar(value='就绪')
        self.count = StringVar(value='已获得 0 条')

        self.build_ui()
        self.on_region_change()
        self.center()

    def build_ui(self):
        P = 18
        ttk.Label(self.root, text='选择地区（先选大区，再选商圈）').grid(
            row=0, column=0, columnspan=5, sticky='w', padx=P, pady=(18, 4))

        self.combo = ttk.Combobox(self.root, textvariable=self.region,
                                  state='readonly', values=REGION_LIST, width=20)
        self.combo.grid(row=1, column=0, columnspan=2, sticky='w', padx=(P, 4))
        self.combo.bind('<<ComboboxSelected>>', self.on_region_change)

        ttk.Label(self.root, text='·').grid(row=1, column=2)

        self.biz_combo = ttk.Combobox(self.root, textvariable=self.biz,
                                      state='readonly', width=16)
        self.biz_combo.grid(row=1, column=3, columnspan=2, sticky='w', padx=(4, P))

        ttk.Label(self.root, text='验证用浏览器', foreground='#666666').grid(
            row=2, column=0, sticky='w', padx=P, pady=(12, 0))
        self.browser_combo = ttk.Combobox(self.root, textvariable=self.browser,
                                          state='readonly', width=20,
                                          values=list(self.browser_map))
        self.browser_combo.grid(row=3, column=0, columnspan=2, sticky='w', padx=(P, 4))

        self.btn_start = ttk.Button(self.root, text='开始采集', command=self.start, width=12)
        self.btn_start.grid(row=4, column=0, columnspan=2, sticky='w',
                            padx=(P, 6), pady=(16, 6))
        self.btn_stop = ttk.Button(self.root, text='停止', command=self.stop,
                                   width=8, state=DISABLED)
        self.btn_stop.grid(row=4, column=3, sticky='w', pady=(16, 6))

        ttk.Separator(self.root, orient='horizontal').grid(
            row=5, column=0, columnspan=5, sticky='ew', padx=P, pady=(10, 0))

        ttk.Label(self.root, textvariable=self.status, foreground='#333333',
                  wraplength=340, justify='left').grid(
            row=6, column=0, columnspan=5, sticky='w', padx=P, pady=(10, 2))
        ttk.Label(self.root, textvariable=self.count, foreground='#1F4E79').grid(
            row=7, column=0, columnspan=5, sticky='w', padx=P, pady=(0, 18))

    def center(self):
        self.root.update_idletasks()
        w, h = self.root.winfo_width(), self.root.winfo_height()
        self.root.geometry(f'+{(self.root.winfo_screenwidth() - w) // 2}'
                           f'+{(self.root.winfo_screenheight() - h) // 3}')

    def on_region_change(self, _=None):
        """大区变了 -> 刷新二级商圈下拉。"""
        region = self.region.get()
        bizs = REGION_BIZ.get(region, [])
        if bizs:
            self.biz_combo.config(state='readonly', values=[BIZ_ALL] + bizs)
        else:                                       # 「全部」没有商圈可选
            self.biz_combo.config(state=DISABLED, values=[BIZ_ALL])
        self.biz.set(BIZ_ALL)

    # ---- 供采集线程回调 ----
    def set_status(self, msg):
        self.root.after(0, lambda: self.status.set(msg))

    def set_count(self, n):
        self.root.after(0, lambda: self.count.set(f'已获得 {n} 条'))

    def notify(self, title, msg):
        self.root.after(0, lambda: messagebox.showwarning(title, msg))

    def ask_verify(self, gate_url, page_url, browser_name, browser_exe, first=True):
        """弹窗询问用户；返回 True 表示继续重试，False 表示停止。"""
        self.verify_done.clear()
        self.verify_ok = False
        self.root.after(0, lambda: self._show_verify(gate_url, page_url,
                                                     browser_name, browser_exe, first))
        while not self.verify_done.is_set():
            time.sleep(0.05)
        return self.verify_ok

    def _show_verify(self, gate_url, page_url, browser_name, browser_exe, first):
        if first:
            open_in_browser(browser_name, browser_exe, gate_url)
            which = browser_name or DEFAULT_BROWSER
            msg = (f'被安居客的人机验证拦住了。\n\n'
                   f'已经用【{which}】打开验证页面，\n'
                   '请按页面提示点一下按钮完成验证，\n'
                   '然后回到本窗口点【确定】继续采集。\n\n'
                   f'如果浏览器没有自动打开，请手动访问：\n{page_url}\n\n')
        else:
            msg = ('仍然被拦截。\n\n'
                   '请确认已经在浏览器里完成了验证码，\n'
                   '再点【确定】重试。\n\n'
                   '如果浏览器里根本没出现验证码，建议先点【取消】停止，'
                   '等 5~10 分钟再试。\n\n')
        self.verify_ok = messagebox.askokcancel(
            '需要人机验证', msg + '点【取消】则停止采集，已采集的数据会保留。')
        self.verify_done.set()

    # ---- 按钮 ----
    def start(self):
        region = self.region.get()
        biz = self.biz.get()
        url, tag = resolve_target(region, biz)

        self.btn_start.config(state=DISABLED)
        self.combo.config(state=DISABLED)
        self.biz_combo.config(state=DISABLED)
        self.browser_combo.config(state=DISABLED)
        self.btn_stop.config(state=NORMAL)
        self.set_count(0)
        self.set_status(f'正在启动…（{tag}）')

        browser_name = self.browser.get()
        self.crawler = Crawler(self, url, tag, browser_name,
                               self.browser_map.get(browser_name, ''))
        self.crawler.start()

    def stop(self):
        self.verify_ok = False
        self.verify_done.set()
        if self.crawler and self.crawler.is_alive():
            self.crawler.stop_event.set()
            self.set_status('正在停止…')

    def finish(self, total, path, stopped):
        def _do():
            self.btn_start.config(state=NORMAL)
            self.combo.config(state='readonly')
            # 只恢复二级下拉的可用状态，保留用户原来选的商圈，方便再采一次
            self.biz_combo.config(
                state='readonly' if REGION_BIZ.get(self.region.get()) else DISABLED)
            self.browser_combo.config(state='readonly')
            self.btn_stop.config(state=DISABLED)
            self.set_count(total)
            self.set_status('已停止' if stopped else '采集完成')
            if not total:
                if stopped:
                    messagebox.showinfo('已停止', '已停止采集，本次没有获取到数据。')
                else:
                    messagebox.showwarning(
                        '没有数据',
                        '这次没有采集到房源。\n\n'
                        '如果浏览器里出现过验证页面，请先在浏览器完成验证，'
                        '然后重新点【开始采集】。\n'
                        '如果反复如此，建议等 5~10 分钟或换个网络再试。')
                return
            head = f'已停止，本次共获取 {total} 条房源。' if stopped else f'共采集 {total} 条房源。'
            messagebox.showinfo(
                '已停止' if stopped else '采集完成',
                f'{head}\n\n'
                f'文件已保存到桌面：\n{os.path.basename(path)}\n\n'
                '打开后可直接点击最后一列「房源链接」跳转到房源页面。')
        self.root.after(0, _do)

    def run(self):
        self.root.mainloop()


def main():
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    App().run()


if __name__ == '__main__':
    main()
